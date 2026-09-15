"""Bounded, local-only BOM extraction. JSON bytes in/out; never executes workbook content."""
import base64
import io
import json
import os
import re
import sys
import zipfile
import pdfplumber

packages = os.environ.get('DLE_OS_SIM_ANALYSIS_PACKAGES')
if packages:
    sys.path.insert(0, packages)


def no_external(event, args):
    if event.startswith('socket.') or event in ('subprocess.Popen', 'os.system', 'os.exec', 'os.posix_spawn'):
        raise RuntimeError('External execution disabled')


sys.addaudithook(no_external)


def norm(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().upper()


def header(value):
    return re.sub(r'[^A-Z0-9]', '', norm(value))


def column(n):
    out = ''
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def text(value, fmt=''):
    if value is None:
        return ''
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if int(value) == value:
            return str(int(value)).zfill(len(fmt)) if re.fullmatch('0+', fmt or '') else str(int(value))
    return str(value).strip()


def spreadsheet(binary):
    if binary.startswith(b'PK'):
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(binary)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Workbook expansion limit')
        book = openpyxl.load_workbook(io.BytesIO(binary), read_only=True, data_only=True, keep_links=False)
        sheets = []
        try:
            if len(book.worksheets) > 20:
                raise ValueError('Sheet limit')
            for sheet in book.worksheets:
                if sheet.max_row > 10000 or sheet.max_column > 100:
                    raise ValueError('Sheet dimensions')
                sheets.append((sheet.title, [[text(c.value, c.number_format) for c in row] for row in sheet.iter_rows()]))
        finally:
            book.close()
        return sheets
    import xlrd
    book = xlrd.open_workbook(file_contents=binary, formatting_info=True, on_demand=True)
    try:
        if book.nsheets > 20:
            raise ValueError('Sheet limit')
        sheets = []
        for sheet in book.sheets():
            if sheet.nrows > 10000 or sheet.ncols > 100:
                raise ValueError('Sheet dimensions')
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    row.append('' if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR) else text(cell.value, fmt))
                rows.append(row)
            sheets.append((sheet.name, rows))
        return sheets
    finally:
        book.release_resources()


def evidence(doc, page=None, sheet=None, location=''):
    return dict(documentId=doc, page=page, sheet=sheet, location=location)


def extract(job, documents):
    governing = next(d for d in documents if d['source']['documentId'] == job['governingDocumentId'])
    rows = []
    pages = set()
    with pdfplumber.open(io.BytesIO(base64.b64decode(governing['binary']))) as pdf:
        if len(pdf.pages) > 50:
            raise ValueError('PDF page limit')
        for page_no, page in enumerate(pdf.pages, 1):
            for table in page.find_tables():
                data = table.extract()
                if not data or [header(v) for v in data[0]] != ['LINE', 'PN', 'QUANTITY', 'DESIGNATOR', 'DESC']:
                    continue
                pages.add(page_no)
                for index, values in enumerate(data[1:], 1):
                    if not any(values):
                        continue
                    if len(values) != 5 or any(v is None for v in values) or not str(values[0]).strip().isdigit():
                        raise ValueError('Incomplete governing row')
                    values = [v.strip() for v in values]
                    fields = {key: dict(value=value, evidence=evidence(job['governingDocumentId'], page_no,
                              location='BOM row ' + str(index) + '; bounds ' + str(table.rows[index].bbox)),
                              uncertainty='Source text preserved; human review required.', relationship='NOT_COMPARED',
                              supportingValue=None, supportingEvidence=None)
                              for key, value in zip(['lineNumber', 'partNumber', 'quantity', 'designators', 'description'], values)}
                    rows.append(dict(fields=fields, manufacturerProposals=[], manufacturerUncertainty='No matching manufacturer identity found.'))
    if not rows or len(rows) > 1000:
        raise ValueError('Governing row limit')
    sources = []
    reports = []
    for doc in documents:
        source = doc['source']
        if source.get('profile', {}).get('derivedAnalysisPurpose') != 'MANUFACTURER_ENRICHMENT':
            continue
        try:
            if source['mimeType'] not in ('application/vnd.ms-excel', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):
                raise ValueError('Unsupported enrichment format')
            tables = 0
            for sheet, cells in spreadsheet(base64.b64decode(doc['binary'])):
                for hr, values in enumerate(cells[:30]):
                    headers = [header(v) for v in values]
                    pn = next((i for i, h in enumerate(headers) if h in ('PARTNO', 'CUSTOMERPN', 'CUSTOMERPARTNUMBER', 'BOMPN', 'PN')), None)
                    mfgs = [i for i, h in enumerate(headers) if h in ('MANUFACTURERPARTNUMBER', 'MFGPN', 'MFGPARTNUMBER') or re.fullmatch(r'ALT\d+MANUFACTURERPARTNUMBER', h)]
                    if pn is None or not mfgs:
                        continue
                    tables += 1
                    def find(names):
                        return next((i for i, h in enumerate(headers) if h in names), None)
                    line = find(('LINO', 'LINE', 'LINENO', 'FIND', 'FINDNO'))
                    assembly = find(('BOMNO', 'ASSEMBLY', 'ASSEMBLYNUMBER'))
                    qty = find(('BOMQTY', 'QUANTITY', 'QTY'))
                    desc = find(('DESCRIPTION', 'DESC'))
                    refs = find(('REFDES', 'DESIGNATORS', 'DESIGNATOR'))
                    for rn, values in enumerate(cells[hr + 1:], hr + 2):
                        def val(i):
                            return values[i] if i is not None and i < len(values) else ''
                        if not val(pn) or (assembly is not None and norm(val(assembly)) != norm(job['assembly'])):
                            continue
                        for mfg in mfgs:
                            if not val(mfg):
                                continue
                            label = headers[mfg]
                            name_header = label.replace('MANUFACTURERPARTNUMBER', 'MFGNAME')
                            name_idx = find((name_header,))
                            if name_idx is None and not label.startswith('ALT'):
                                name_idx = find(('MFGNAME', 'MANUFACTURER', 'MANUFACTURERNAME'))
                            sources.append(dict(customer=val(pn), line=val(line), quantity=val(qty), description=val(desc), designators=val(refs),
                                manufacturer=val(name_idx) or None, partNumber=val(mfg), sourceLabel=values and cells[hr][mfg],
                                evidence=evidence(source['documentId'], sheet=sheet, location=column(mfg + 1) + str(rn)),
                                customerEvidence=evidence(source['documentId'], sheet=sheet, location=column(pn + 1) + str(rn))))
                    break
            reports.append(source['documentId'] + (': parsed locally; ' + str(tables) + ' recognized table(s).' if tables else ': NOT PROCESSED — no recognized customer/MFG header pair.'))
        except ImportError:
            reports.append(source['documentId'] + ': NOT PROCESSED — local workbook reader dependency unavailable.')
        except Exception:
            reports.append(source['documentId'] + ': NOT PROCESSED — unsupported/unreadable workbook or safety limit; no identities inferred.')
    for row_index, row in enumerate(rows):
        fields = row['fields']
        matches = [s for s in sources if norm(s['customer']) == norm(fields['partNumber']['value'])]
        exact_line = [s for s in matches if s['line'] and norm(s['line']) == norm(fields['lineNumber']['value'])]
        if exact_line:
            matches = exact_line
        for n, s in enumerate(matches):
            basis = ['Exact customer/BOM P/N']
            conflicts = []
            if s['line'] == fields['lineNumber']['value']:
                basis.append('Exact Find / line number')
            elif s['line']:
                conflicts.append('Source line number differs')
            for key in ('quantity', 'description', 'designators'):
                if not s[key] or norm(s[key]) in ('SEE DRAWING', '--'):
                    continue
                if norm(s[key]) == norm(fields[key]['value']):
                    basis.append(key + ' agrees')
                else:
                    conflicts.append(key + ' differs; governing value retained')
            row['manufacturerProposals'].append(dict(id='mfg-' + str(row_index + 1) + '-' + str(n + 1),
                manufacturerName=s['manufacturer'], partNumber=s['partNumber'], evidence=s['evidence'], customerEvidence=s['customerEvidence'],
                governingEvidence=fields['partNumber']['evidence'], sourceLabel=s['sourceLabel'], matchBasis=basis, conflicts=conflicts,
                sourceValues={k:s[k] for k in ('customer', 'line', 'quantity', 'description', 'designators')},
                confidence='HIGH' if exact_line and not conflicts else 'MEDIUM',
                uncertainty='Multiple source candidates; human selection required.' if len(matches) > 1 else 'Human identity confirmation required.'))
        if matches:
            row['manufacturerUncertainty'] = 'Multiple candidates require review.' if len(matches) > 1 else 'Proposed from source evidence; not approved for substitution.'
        elif any('NOT PROCESSED' in report for report in reports):
            row['manufacturerUncertainty'] = 'Unresolved: one or more enrichment sources were not processed. See analysis coverage.'
    return dict(contractVersion='DLE_CANDIDATE_ANALYSIS_RESULT_V2', outcome='EXTRACTED', coverage='TABLES_COMPLETE',
        coverageReason=str(len(rows)) + ' governing rows: all rows of recognized BOM tables on PDF page(s) ' + ', '.join(map(str, sorted(pages))) +
        '. Other drawing content was not interpreted; no OCR. ' + ' '.join(reports), rows=rows)


if __name__ == '__main__':
    try:
        payload = json.loads(sys.stdin.buffer.read(115 * 1024 * 1024))
        print(json.dumps(extract(payload['job'], payload['documents']), ensure_ascii=True))
    except Exception:
        sys.exit(2)
